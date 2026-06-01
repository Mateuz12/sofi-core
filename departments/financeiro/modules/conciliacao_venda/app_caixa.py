import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
import openpyxl
import re
import os

# --- 1. CONFIGURAÇÃO FIXA ---
# Coloque o caminho absoluto da sua planilha de controle AQUI (ela não muda)
CAMINHO_CONTROLE = r"\\srv-dc01\File Server\05-FIN\FECHAMENTO DE CAIXA\CONTROLE DE DEPÓSITO\UDI\2026\05-FIN.FECHAMENTO DE CAIXA.CONTROLE DE DEPÓSITO.UDI.2026.PLANILHA DE ACOMPANHAMENTO DE DEPÓSITOS LOJAS - VD.xlsx"

# --- 2. O MOTOR DO APLICATIVO (Função principal) ---
def processar_automacao():
    caminho_retaguarda = lbl_caminho.cget("text")
    
    # Trava 1: Usuário esqueceu de escolher o arquivo
    if caminho_retaguarda == "Nenhum arquivo selecionado":
        messagebox.showwarning("Aviso", "Por favor, selecione o relatório do Retaguarda primeiro.")
        return

    # Trava 2: Verifica se a planilha de controle está fechada
    try:
        with open(CAMINHO_CONTROLE, 'a'):
            pass
    except IOError:
        messagebox.showerror("Erro Crítico", "A sua Planilha de Controle está ABERTA no Excel.\nFeche a planilha e tente novamente.")
        return

    try:
        lbl_status.config(text="Status: Lendo relatório e calculando...", fg="#76b5c5")
        janela.update() # Força a tela a atualizar o texto

        # -- SEU CÓDIGO ETL ENTRA AQUI --
        
        # 1. Lê a data na F2
        df_filtro = pd.read_excel(caminho_retaguarda, usecols="F", header=None, nrows=2)
        texto_f2 = str(df_filtro.iloc[1, 0]) 
        busca_data = re.search(r'(\d{2}/\d{2}/\d{4})', texto_f2)
        
        if not busca_data:
            messagebox.showerror("Erro de Leitura", "Não foi possível encontrar a data na célula F2 do relatório.")
            lbl_status.config(text="Status: Aguardando ação.", fg="white")
            return
            
        data_extraida = busca_data.group(1)
        aba_alvo = data_extraida.replace('/', '.')
        
        # 2. Extrai os valores usando sua regra de Cascata (Forward Fill)
        df_retaguarda = pd.read_excel(caminho_retaguarda, usecols="B,X", names=["Descricao", "Valor"])
        df_retaguarda['Codigo'] = df_retaguarda['Descricao'].astype(str).str.extract(r'Loja:\s*(\d+)')
        df_retaguarda['Codigo'] = df_retaguarda['Codigo'].ffill()
        df_dinheiro = df_retaguarda[df_retaguarda['Descricao'].astype(str).str.contains('001 - DINHEIRO', na=False)].copy()

        # Função de limpar moeda
        def limpar_moeda(valor):
            if pd.isna(valor): return 0.0
            valor_str = str(valor).strip().replace('.', '').replace(',', '.')
            try: return float(valor_str)
            except ValueError: return 0.0

        df_dinheiro['Valor'] = df_dinheiro['Valor'].apply(limpar_moeda)
        df_limpo = df_dinheiro.dropna(subset=['Codigo', 'Valor'])
        dicionario_valores = dict(zip(df_limpo['Codigo'], df_limpo['Valor']))

        # 3. Injeta na Planilha de Controle
        wb = openpyxl.load_workbook(CAMINHO_CONTROLE)
        
        if aba_alvo not in wb.sheetnames:
            messagebox.showerror("Erro de Aba", f"A aba '{aba_alvo}' não existe na planilha de controle.")
            lbl_status.config(text="Status: Aguardando ação.", fg="white")
            return

        planilha = wb[aba_alvo]
        lojas_atualizadas = 0

        for linha in range(2, planilha.max_row + 1):
            celula_codigo = planilha.cell(row=linha, column=2).value 
            if celula_codigo:
                codigo_str = str(celula_codigo).strip() 
                if codigo_str in dicionario_valores:
                    planilha.cell(row=linha, column=6).value = dicionario_valores[codigo_str]
                    lojas_atualizadas += 1

        wb.save(CAMINHO_CONTROLE)
        
        # Alerta de Sucesso
        lbl_status.config(text=f"Concluído! {lojas_atualizadas} lojas atualizadas.", fg="#8ce071")
        messagebox.showinfo("Sucesso Absoluto", f"Planilha atualizada com sucesso!\n\nData processada: {aba_alvo}\nLojas atualizadas: {lojas_atualizadas}")

    except Exception as e:
        # Se QUALQUER coisa der errado, ele avisa sem fechar o app
        lbl_status.config(text="Status: Falha na execução.", fg="#ff6b6b")
        messagebox.showerror("Erro Inesperado", f"Ocorreu um erro no código:\n\n{str(e)}")

# --- 3. FUNÇÃO DO BOTÃO DE SELEÇÃO ---
def selecionar_arquivo():
    caminho = filedialog.askopenfilename(
        title="Selecione o relatório do Retaguarda",
        filetypes=[("Arquivos Excel", "*.xls *.xlsx")]
    )
    if caminho:
        lbl_caminho.config(text=caminho)
        lbl_status.config(text="Status: Arquivo carregado. Pronto para rodar.", fg="white")

# --- 4. CONSTRUÇÃO DA INTERFACE (A Tela) ---
janela = tk.Tk()
janela.title("Conciliador de Fundo de Caixa - Finanças GB")
janela.geometry("550x250")

# Definir a cor que mais se aproxima do fundo da imagem para manter a harmonia (verde escuro da GB)
COR_FUNDO = "#022111" 
COR_TEXTO = "white"
janela.configure(bg=COR_FUNDO)

# Adicionar imagem de fundo
try:
    from PIL import Image, ImageTk
    # Tenta carregar a imagem e redimensionar para o tamanho da janela
    imagem_original = Image.open("fundo.png")
    try:
        resampling_filter = Image.Resampling.LANCZOS
    except AttributeError:
        resampling_filter = Image.LANCZOS # Suporte para versoes antigas do Pillow
        
    imagem_red = imagem_original.resize((550, 250), resampling_filter)
    img_fundo = ImageTk.PhotoImage(imagem_red)
    
    lbl_fundo = tk.Label(janela, image=img_fundo, borderwidth=0, highlightthickness=0, bg=COR_FUNDO)
    lbl_fundo.place(x=0, y=0, relwidth=1, relheight=1)
except Exception as e:
    print(f"Aviso: Não foi possível carregar ou redimensionar a imagem de fundo usando PIL: {e}")
    # Fallback usando tk.PhotoImage puro caso Pillow falhe
    try:
        img_fundo = tk.PhotoImage(file="fundo.png")
        lbl_fundo = tk.Label(janela, image=img_fundo, borderwidth=0, highlightthickness=0, bg=COR_FUNDO)
        lbl_fundo.place(x=0, y=0, relwidth=1, relheight=1)
    except Exception as erro_fallback:
        print(f"Aviso: Fallback tk.PhotoImage falhou. Certifique-se de que a imagem seja um arquivo válido 'fundo.png'.")

# Usar um Frame com cor de fundo que combine para agrupar os itens
frame_main = tk.Frame(janela, bg=COR_FUNDO)
frame_main.pack(fill="both", expand=True, padx=20, pady=20)

# Textos
titulo = tk.Label(frame_main, text="Automação de Fechamento", font=("Arial", 14, "bold"), bg=COR_FUNDO, fg=COR_TEXTO)
titulo.pack(pady=5)

instrucao = tk.Label(frame_main, text="1. Selecione o relatório do Retaguarda baixado hoje:", font=("Arial", 10), bg=COR_FUNDO, fg=COR_TEXTO)
instrucao.pack(anchor="w")

# Botão de procurar e texto do caminho
frame_arquivo = tk.Frame(frame_main, bg=COR_FUNDO)
frame_arquivo.pack(fill="x", pady=5)

btn_procurar = tk.Button(frame_arquivo, text="Buscar Arquivo", command=selecionar_arquivo, bg="#e0e0e0", fg="black")
btn_procurar.pack(side="left")

lbl_caminho = tk.Label(frame_arquivo, text="Nenhum arquivo selecionado", fg="#cccccc", bg=COR_FUNDO, font=("Arial", 9, "italic"))
lbl_caminho.pack(side="left", padx=10)

# O Botão Principal
btn_rodar = tk.Button(frame_main, text="⚙️ INJETAR DADOS NA PLANILHA", command=processar_automacao, bg="#4CAF50", fg="white", font=("Arial", 12, "bold"), pady=10)
btn_rodar.pack(fill="x", pady=20)

# Barra de status no rodapé
lbl_status = tk.Label(frame_main, text="Status: Aguardando ação.", font=("Arial", 9), bg=COR_FUNDO, fg=COR_TEXTO)
lbl_status.pack(side="bottom")

# Inicia o aplicativo
janela.mainloop()