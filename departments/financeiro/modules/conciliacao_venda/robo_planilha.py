import pandas as pd
import openpyxl
import re

print("1. Lendo o relatório do Retaguarda...")
caminho_retaguarda = r"C:\Users\luma.dias\Downloads\Relatorio_FormasDePagamentoAnalitico (11).xls"

#  DESCOBRINDO A DATA 
print("-> Procurando a data na célula F2...")
# Ler apenas a coluna F (índice 5 no Excel) e as duas primeiras linhas
df_filtro = pd.read_excel(caminho_retaguarda, usecols="F", header=None, nrows=2)

# Pega o texto exato da célula F2 (Linha 2 é o índice 1 no Pandas)
texto_f2 = str(df_filtro.iloc[1, 0]) 

# Usa Regex para capturar o padrão de data"
busca_data = re.search(r'(\d{2}/\d{2}/\d{4})', texto_f2)

if busca_data:
    data_extraida = busca_data.group(1) # Pega o "05/03/2026"
    aba_alvo = data_extraida.replace('/', '.') # Transforma em "05.03.2026"
    print(f"-> Data identificada com sucesso: Aba alvo será '{aba_alvo}'")

# Se não conseguir ler a data na planilha automaticamente, pedir para colocar a data manualmente. 

else:
    print("-> AVISO: Não consegui ler a data na célula F2.")
    aba_alvo = input('Por segurança, digite a data manualmente (Ex: 05.03.2026): ') 

# ETAPA DE EXTRAÇÃO DOS VALOLES
df_retaguarda = pd.read_excel(caminho_retaguarda, usecols="B,X", names=["Descricao", "Valor"])

# Extrair código da loja(Forward Fill)
df_retaguarda['Codigo'] = df_retaguarda['Descricao'].astype(str).str.extract(r'Loja:\s*(\d+)')
df_retaguarda['Codigo'] = df_retaguarda['Codigo'].ffill()

# Filtra APENAS o dinheiro em espécie e usa .copy() para evitar avisos do Pandas
df_dinheiro = df_retaguarda[df_retaguarda['Descricao'].astype(str).str.contains('001 - DINHEIRO', na=False)].copy()

# Tranformando o texto em numérico e substituindo a vírgula por pontos
def limpar_moeda(valor):
    if pd.isna(valor):
        return 0.0
    valor_str = str(valor).strip()
    
    valor_str = valor_str.replace('.', '')
    
    valor_str = valor_str.replace(',', '.')
    try:
        return float(valor_str) # Converte para número matemático real
    except ValueError:
        return 0.0

# Aplica a limpeza na coluna de Valor
df_dinheiro['Valor'] = df_dinheiro['Valor'].apply(limpar_moeda)

# Limpa e cria o Dicionário
df_limpo = df_dinheiro.dropna(subset=['Codigo', 'Valor'])
dicionario_valores = dict(zip(df_limpo['Codigo'], df_limpo['Valor']))

print(f"-> {len(dicionario_valores)} lojas prontas para injeção.")


# Preenchimento da planilha de controle 
print(f"\n2. Abrindo a Planilha de Controle na aba {aba_alvo}...")
caminho_controle = r"C:\Users\luma.dias\Desktop\depositos.xlsx"
wb = openpyxl.load_workbook(caminho_controle)

if aba_alvo not in wb.sheetnames:
    print(f"\nERRO: A aba '{aba_alvo}' não existe na sua planilha de controle.")
    exit()

planilha = wb[aba_alvo]
lojas_atualizadas = 0

for linha in range(2, planilha.max_row + 1):
    celula_codigo = planilha.cell(row=linha, column=2).value 
    
    if celula_codigo:
        codigo_str = str(celula_codigo).strip() 
        
        if codigo_str in dicionario_valores:
            planilha.cell(row=linha, column=6).value = dicionario_valores[codigo_str]
            lojas_atualizadas += 1

wb.save(caminho_controle)
print(f"\nSUCESSO FINAL! {lojas_atualizadas} lojas foram atualizadas automaticamente. Fórmulas preservadas.")