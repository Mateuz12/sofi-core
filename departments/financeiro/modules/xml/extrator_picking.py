"""
extrator_picking.py - Automação semanal para extração de OP Picking via API do Grupo Boticário.

Estrutura da Branch (T:\\13-Publica):
- Pasta base de relatórios:
    T:\\13-Publica\\XMLS CÁLAMO\\RELATÓRIOS COMPLETOS\\<Ano>\\<Mês>\\<Semana>\\
    (com os relatórios: Relatorio_Final_Matheus_BH.xls e Relatorio_Final_Matheus_UDI.xls)
- Pasta base de XMLs organizados:
    T:\\13-Publica\\XMLS CÁLAMO\\XML BAIXADOS\\<Ano>\\<Mês>\\<Semana>\\<PDV>\\
    (com os arquivos XML nomeados como: <PDV>_<Nº_DOC>.xml)
- Pasta e nomes das planilhas de Picking geradas:
    Salvas na mesma subpasta do PDV com o mesmo nome do XML:
    T:\\13-Publica\\XMLS CÁLAMO\\XML BAIXADOS\\<Ano>\\<Mês>\\<Semana>\\<PDV>\\<PDV>_<Nº_DOC>.xlsx

Fluxo:
1. Localiza a ÚLTIMA semana gerada dos relatórios finais do Matheus em T:\\13-Publica (ou local).
2. Extrai todas as notas faturadas: numero_nota, nosso_pedido, pdv, nome_loja, data_emissao, unidade.
3. Localiza a pasta correspondente de XML BAIXADOS da mesma semana e mapeia as subpastas de PDV.
4. Efetua login no portal Extranet via Playwright e captura os tokens de sessão (Bearer JWT).
5. Para cada nota e PDV, consulta a API de picking com rate-limit (delay >= 5s com jitter).
6. Gera uma planilha individual (.xlsx) no padrão exato de op_piking.xlsx (aba 'data', 9 colunas)
   salva em: <pasta_xmls_semana>/<PDV>/<PDV>_<numero_nota>.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import random
import re
import sys
import time
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import requests
import xlrd
from playwright.sync_api import (
    BrowserContext,
    Locator,
    Page,
    sync_playwright,
)

# ─────────────────────────────────────────────
# Localização de Raiz e Carregamento de Ambiente
# ─────────────────────────────────────────────

_ROOT = Path(__file__).resolve().parent
while not (_ROOT / ".env").exists() and _ROOT != _ROOT.parent:
    _ROOT = _ROOT.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

try:
    from core.sofi_env import load_sofi_env
    load_sofi_env(department="financeiro")
except Exception as _env_err:
    print(f"[Aviso SOFI Env] {_env_err}")

try:
    from core.database.logger import log_execution
except Exception:
    def log_execution(modulo="XML Picking Semanal"):
        def decorator(func):
            def wrapper(*args, **kwargs):
                return func(*args, **kwargs)
            return wrapper
        return decorator

# ─────────────────────────────────────────────
# Constantes e Configurações (Branch T:\13-Publica)
# ─────────────────────────────────────────────

DIRETORIO_ATUAL = Path(__file__).resolve().parent
API_BASE_URL = "https://list-api.prd.demanda-abastecimento.grupoboticario.digital"

# Caminhos base da rede T:\13-Publica
REDE_BASE_T = Path(r"T:\13-Publica\XMLS CÁLAMO")
REDE_RELATORIOS_BASE = REDE_BASE_T / "RELATÓRIOS COMPLETOS"
REDE_XMLS_BASE = REDE_BASE_T / "XML BAIXADOS"

COLUNAS_OP_PIKING = [
    "volume",
    "itemNumber",
    "handlingUnitCode",
    "quantity",
    "batchCode",
    "volumeIdentifier",
    "materialCode",
    "materialDescription",
    "expirationDate",
]

XPATHS = {
    "user": '//*[@id="signInName"]',
    "password": '//*[@id="password"]',
}

REALISTIC_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

STEALTH_INIT_SCRIPT = """
Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
Object.defineProperty(navigator, 'plugins', { get: () => [1, 2, 3, 4, 5] });
Object.defineProperty(navigator, 'languages', { get: () => ['pt-BR', 'pt', 'en-US', 'en'] });
"""

# ─────────────────────────────────────────────
# Funções Utilitárias de Ambiente
# ─────────────────────────────────────────────

def env_bool(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "sim", "s"}


def required_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError(f"Preencha a variável {name} no arquivo .env antes de executar.")
    return value


def browser_args() -> list[str]:
    args = [
        "--disable-quic",
        "--ignore-certificate-errors",
        "--disable-blink-features=AutomationControlled",
        "--disable-features=IsolateOrigins,site-per-process",
        "--no-first-run",
        "--no-default-browser-check",
    ]
    if env_bool("DISABLE_HTTP2", default=True):
        args.append("--disable-http2")
    return args


def by_xpath(page_or_locator: Page | Locator, xpath: str) -> Locator:
    return page_or_locator.locator(f"xpath={xpath}")


# ─────────────────────────────────────────────
# 1. Leitura dos Relatórios do Matheus
# ─────────────────────────────────────────────

def _limpar_valor_texto(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return str(val)
    return str(val).strip()


def extrair_notas_relatorio_matheus(caminho_arquivo: Path | str) -> List[Dict[str, Any]]:
    """
    Lê com precisão as planilhas do relatório final do Matheus (BH e UDI)
    e extrai os registros das notas:
    - numero_nota (Nº NF)
    - nosso_pedido (Nosso Pedido)
    - pdv (Código da Loja)
    - nome_loja
    - data_emissao
    - unidade (BH ou UDI)
    """
    caminho = Path(caminho_arquivo)
    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    registros: List[Dict[str, Any]] = []
    chaves_vistas = set()

    is_udi = "UDI" in caminho.name.upper()
    unidade = "UDI" if is_udi else "BH"

    print(f"Lendo notas do relatório do Matheus ({unidade}): {caminho.name}...")

    # Formato .XLS (Padrão dos relatórios do Matheus salvos pelo Excel COM)
    if caminho.suffix.lower() == ".xls":
        wb = xlrd.open_workbook(caminho)
        ws_base = wb.sheet_by_name("Base") if "Base" in wb.sheet_names() else None

        if "Notas Fiscais" in wb.sheet_names():
            ws = wb.sheet_by_name("Notas Fiscais")
            
            # BH:  Col 0: PDV | Col 1: Nome | Col 2: Dt. Emissão | Col 3: Nº NF | Col 4: Nosso Pedido
            # UDI: Col 5: PDV | Col 6: Nome | Col 7: Dt. Emissão | Col 8: Nº NF | Col 9: Nosso Pedido
            col_pdv = 5 if is_udi else 0
            col_nome = 6 if is_udi else 1
            col_emi = 7 if is_udi else 2
            col_nf = 8 if is_udi else 3
            col_ped = 9 if is_udi else 4

            for r in range(3, ws.nrows):
                row = ws.row_values(r)
                if len(row) <= max(col_pdv, col_nf, col_ped):
                    continue

                pdv_raw = row[col_pdv]
                nf_raw = row[col_nf]
                ped_raw = row[col_ped]
                nome_raw = row[col_nome] if len(row) > col_nome else ""
                emi_raw = row[col_emi] if len(row) > col_emi else ""

                row_base_idx = r - 2 # Row 3 de Notas Fiscais mapeia para Row 1 de Base (após cabeçalho)
                
                # Se NF ou PDV estiverem como fórmula avaliada em 0 pelo xlrd, obtém da Base
                if (not nf_raw or nf_raw == 0.0 or str(nf_raw).strip() == "") and ws_base and row_base_idx < ws_base.nrows:
                    nf_raw = ws_base.cell_value(row_base_idx, 3) # Coluna D = Número do Documento
                
                if (not pdv_raw or pdv_raw == 0.0 or str(pdv_raw).strip() == "") and ws_base and row_base_idx < ws_base.nrows:
                    pdv_raw = ws_base.cell_value(row_base_idx, 0) # Coluna A = Código da Loja

                pdv_str = _limpar_valor_texto(pdv_raw)
                nf_str = _limpar_valor_texto(nf_raw)
                ped_str = _limpar_valor_texto(ped_raw)
                nome_str = _limpar_valor_texto(nome_raw)
                emi_str = _limpar_valor_texto(emi_raw)

                # Se a data veio em formato de serial do Excel
                if emi_raw and isinstance(emi_raw, float) and emi_raw > 30000:
                    try:
                        dt_tuple = xlrd.xldate_as_tuple(emi_raw, wb.datemode)
                        emi_str = f"{dt_tuple[2]:02d}/{dt_tuple[1]:02d}/{dt_tuple[0]}"
                    except Exception:
                        pass

                if pdv_str and nf_str and pdv_str != "0" and nf_str != "0":
                    chave = (nf_str, pdv_str)
                    if chave not in chaves_vistas:
                        chaves_vistas.add(chave)
                        registros.append({
                            "numero_nota": nf_str,
                            "nosso_pedido": ped_str,
                            "pdv": pdv_str,
                            "nome_loja": nome_str,
                            "data_emissao": emi_str,
                            "unidade": unidade,
                        })

    # Formato .XLSX
    elif caminho.suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(caminho, data_only=True)
        if "Notas Fiscais" in wb.sheetnames:
            ws = wb["Notas Fiscais"]
            col_pdv = 6 if is_udi else 1
            col_nome = 7 if is_udi else 2
            col_emi = 8 if is_udi else 3
            col_nf = 9 if is_udi else 4
            col_ped = 10 if is_udi else 5

            for r in range(4, ws.max_row + 1):
                pdv_str = _limpar_valor_texto(ws.cell(r, col_pdv).value)
                nome_str = _limpar_valor_texto(ws.cell(r, col_nome).value)
                emi_str = _limpar_valor_texto(ws.cell(r, col_emi).value)
                nf_str = _limpar_valor_texto(ws.cell(r, col_nf).value)
                ped_str = _limpar_valor_texto(ws.cell(r, col_ped).value)

                if pdv_str and nf_str and pdv_str != "0" and nf_str != "0":
                    chave = (nf_str, pdv_str)
                    if chave not in chaves_vistas:
                        chaves_vistas.add(chave)
                        registros.append({
                            "numero_nota": nf_str,
                            "nosso_pedido": ped_str,
                            "pdv": pdv_str,
                            "nome_loja": nome_str,
                            "data_emissao": emi_str,
                            "unidade": unidade,
                        })

    print(f"  -> {len(registros)} notas extraídas com sucesso de {caminho.name}")
    return registros


def localizar_ultimos_relatorios_matheus() -> Tuple[List[Path], Optional[Path]]:
    """
    Busca automaticamente os relatórios finais do Matheus (Relatorio_Final_Matheus_BH e UDI)
    da ÚLTIMA semana gerada na rede corporativa T:\13-Publica ou localmente.
    Retorna a lista de relatórios e o Path da pasta da semana encontrada.
    """
    relatorios: List[Path] = []
    pasta_semana_encontrada: Optional[Path] = None
    
    # 1. Busca na rede T:\13-Publica\XMLS CÁLAMO\RELATÓRIOS COMPLETOS
    if REDE_RELATORIOS_BASE.exists():
        pastas_semana = [p for p in REDE_RELATORIOS_BASE.glob("*/*/*") if p.is_dir()]
        semanas_com_matheus = [
            p for p in pastas_semana
            if any("matheus" in f.name.lower() for f in p.glob("*.xls*"))
        ]
        if semanas_com_matheus:
            pasta_semana_encontrada = max(semanas_com_matheus, key=lambda p: p.stat().st_mtime)
            print(f"\n[OK] Última semana de relatórios do Matheus detectada em T:\\13-Publica: {pasta_semana_encontrada}")
            for f in pasta_semana_encontrada.glob("*.xls*"):
                if not f.name.startswith("~$") and "matheus" in f.name.lower():
                    relatorios.append(f)

    # 2. Se não encontrou em 13-Publica, tenta o fallback local
    if not relatorios:
        pasta_local_rel = DIRETORIO_ATUAL / "Relatórios Cálamo"
        if pasta_local_rel.exists():
            pastas_loc = [p for p in pasta_local_rel.glob("*/*/*") if p.is_dir()]
            semanas_loc_matheus = [
                p for p in pastas_loc
                if any("matheus" in f.name.lower() for f in p.glob("*.xls*"))
            ]
            if semanas_loc_matheus:
                pasta_semana_encontrada = max(semanas_loc_matheus, key=lambda p: p.stat().st_mtime)
                print(f"\n[OK] Última semana local detectada: {pasta_semana_encontrada}")
                for f in pasta_semana_encontrada.glob("*.xls*"):
                    if not f.name.startswith("~$") and "matheus" in f.name.lower():
                        relatorios.append(f)

    # 3. Fallback: arquivos soltos na pasta do módulo
    if not relatorios:
        for f in DIRETORIO_ATUAL.glob("*Matheus*.xls*"):
            if not f.name.startswith("~$") and "modelo" not in f.name.lower():
                relatorios.append(f)
        pasta_semana_encontrada = DIRETORIO_ATUAL

    return relatorios, pasta_semana_encontrada


# ─────────────────────────────────────────────
# 2. Localização da Pasta de XMLs por PDV
# ─────────────────────────────────────────────

def localizar_pasta_xmls_semana(pasta_relatorio: Optional[Path] = None) -> Path:
    r"""
    Localiza a pasta correspondente de XMLs baixados para a semana dos relatórios.
    Exemplo: T:\13-Publica\XMLS CÁLAMO\XML BAIXADOS\<Ano>\<Mês>\<Semana>
    """
    if pasta_relatorio:
        caminho_str = str(pasta_relatorio)
        if "RELATÓRIOS COMPLETOS" in caminho_str.upper() or "RELATORIOS COMPLETOS" in caminho_str.upper():
            pasta_xml_candidata = Path(
                re.sub(r"RELAT[OÓ]RIOS COMPLETOS", "XML BAIXADOS", caminho_str, flags=re.IGNORECASE)
            )
            if pasta_xml_candidata.exists():
                return pasta_xml_candidata

    # Busca padrão na rede T:\13-Publica\XMLS CÁLAMO\XML BAIXADOS
    if REDE_XMLS_BASE.exists():
        subpastas = [p for p in REDE_XMLS_BASE.glob("*/*/*") if p.is_dir()]
        if subpastas:
            return max(subpastas, key=lambda p: p.stat().st_mtime)
        return REDE_XMLS_BASE

    # Fallbacks locais
    pasta_local_xmls = DIRETORIO_ATUAL / "XML BAIXADOS"
    if pasta_local_xmls.exists():
        subpastas_loc = [p for p in pasta_local_xmls.glob("*/*/*") if p.is_dir()]
        if subpastas_loc:
            return max(subpastas_loc, key=lambda p: p.stat().st_mtime)
        return pasta_local_xmls

    pasta_meudanfe = DIRETORIO_ATUAL / "xmls_meudanfe"
    if pasta_meudanfe.exists():
        return pasta_meudanfe

    return DIRETORIO_ATUAL


def obter_caminho_planilha_picking(
    pasta_xmls_semana: Path,
    pdv: str,
    numero_nota: str,
) -> Path:
    """
    Determina o caminho de destino da planilha de picking da nota:
    1. Procura na subpasta do PDV (<pasta_semana>/<PDV>/) se já existe o XML
       com o nome padrão <PDV>_<Nº_DOC>.xml ou NFE-...xml
    2. Se existir, adota o mesmo stem + '.xlsx'
    3. Se não existir ainda, garante a criação da subpasta do PDV e define:
       <pasta_semana>/<PDV>/<PDV>_<numero_nota>.xlsx
    """
    subpasta_pdv = pdv if pdv else "PDV_SEM_IDENTIFICACAO"
    pasta_destino_pdv = pasta_xmls_semana / subpasta_pdv
    pasta_destino_pdv.mkdir(parents=True, exist_ok=True)

    # 1. Verifica se já existe um arquivo XML nesta pasta com o número da nota
    nome_padrao = f"{pdv}_{numero_nota}.xlsx"
    if pasta_destino_pdv.exists():
        for f_xml in pasta_destino_pdv.glob("*.xml"):
            # Caso 1: Nome padrão <PDV>_<Nº_DOC>.xml
            if f_xml.stem == f"{pdv}_{numero_nota}":
                return pasta_destino_pdv / f"{f_xml.stem}.xlsx"
            # Caso 2: Chave de NFe que contenha o número da nota
            match_chave = re.search(r"55\d{3}(\d{9})", f_xml.name)
            if match_chave and str(int(match_chave.group(1))) == str(int(numero_nota)):
                return pasta_destino_pdv / f"{f_xml.stem}.xlsx"

    # 2. Padrão da branch: <PDV>_<Nº_DOC>.xlsx dentro da subpasta do PDV
    return pasta_destino_pdv / nome_padrao


# ─────────────────────────────────────────────
# 3. Autenticação Extranet e Captura de Tokens
# ─────────────────────────────────────────────

def _tratar_popup_avaliacao(page: Page) -> None:
    """Verifica se há popup de avaliação (Maze widget) e fecha se houver."""
    try:
        if page.locator("#maze-contextual-widget-host").first.is_visible(timeout=500):
            page.evaluate("""() => {
                var host = document.querySelector('#maze-contextual-widget-host');
                if (host && host.shadowRoot) {
                    var btn = host.shadowRoot.querySelector('button[aria-label*="close" i], button[aria-label*="fechar" i], [class*="close" i]');
                    if (btn) btn.click();
                }
            }""")
    except Exception:
        pass


def autenticar_extranet_e_obter_sessao() -> Dict[str, Any]:
    """
    Efetua login no Extranet do Boticário e captura tokens de autorização (Bearer)
    e cookies de sessão para comunicação com a API.
    """
    url = os.getenv("EXTRANET_URL", "https://extranet.grupoboticario.com.br/")
    user = required_env("EXTRANET_USER")
    password = required_env("EXTRANET_PASSWORD")
    headless = env_bool("HEADLESS", default=False)
    browser_channel = os.getenv("BROWSER_CHANNEL", "chrome")

    print("\n--- [Extranet] Iniciando autenticação no portal ---")
    
    auth_headers: Dict[str, str] = {
        "User-Agent": REALISTIC_USER_AGENT,
        "Accept": "application/json, text/plain, */*",
        "Origin": "https://extranet.grupoboticario.com.br",
        "Referer": "https://extranet.grupoboticario.com.br/",
    }
    bearer_token: Optional[str] = None
    captured_cookies: List[dict] = []

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=headless,
            slow_mo=200,
            channel=browser_channel,
            args=browser_args(),
        )
        context: BrowserContext = browser.new_context(
            accept_downloads=True,
            ignore_https_errors=True,
            user_agent=REALISTIC_USER_AGENT,
            viewport={"width": 1366, "height": 768},
            locale="pt-BR",
        )
        page: Page = context.new_page()
        page.add_init_script(STEALTH_INIT_SCRIPT)
        page.set_default_timeout(45_000)

        def on_request(req):
            nonlocal bearer_token
            auth_val = req.headers.get("authorization") or req.headers.get("Authorization")
            if auth_val and "bearer " in auth_val.lower():
                token = auth_val.split(" ", 1)[1].strip()
                if token and len(token) > 20 and not bearer_token:
                    bearer_token = token
                    print("  [OK] Token de autorização capturado via tráfego de rede!")

        page.on("request", on_request)

        try:
            print("Navegando para o portal Extranet...")
            page.goto(url, wait_until="networkidle", timeout=60_000)
            
            by_xpath(page, XPATHS["user"]).wait_for(state="visible", timeout=60_000)
            by_xpath(page, XPATHS["user"]).fill(user)
            by_xpath(page, XPATHS["password"]).fill(password)
            
            print("Credenciais preenchidas. Submetendo login...")
            time.sleep(2)
            page.keyboard.press("Enter")

            print("Aguardando confirmação de sessão...")
            start_wait = time.time()
            while time.time() - start_wait < 60:
                _tratar_popup_avaliacao(page)
                if bearer_token:
                    break
                page.wait_for_timeout(1_500)

            if not bearer_token:
                try:
                    storage_tokens = page.evaluate("""() => {
                        let tokens = [];
                        for (let i = 0; i < localStorage.length; i++) {
                            let k = localStorage.key(i);
                            let v = localStorage.getItem(k);
                            if (v && (v.includes('eyJ') || k.toLowerCase().includes('token') || k.toLowerCase().includes('auth'))) {
                                tokens.push(v);
                            }
                        }
                        for (let i = 0; i < sessionStorage.length; i++) {
                            let k = sessionStorage.key(i);
                            let v = sessionStorage.getItem(k);
                            if (v && (v.includes('eyJ') || k.toLowerCase().includes('token') || k.toLowerCase().includes('auth'))) {
                                tokens.push(v);
                            }
                        }
                        return tokens;
                    }""")
                    for t in storage_tokens:
                        if isinstance(t, str):
                            try:
                                parsed = json.loads(t)
                                if isinstance(parsed, dict):
                                    cand = parsed.get("accessToken") or parsed.get("access_token") or parsed.get("idToken") or parsed.get("secret")
                                    if cand and isinstance(cand, str) and cand.startswith("eyJ"):
                                        bearer_token = cand
                                        print("  [OK] Token de autorização extraído do Storage do navegador!")
                                        break
                            except Exception:
                                if t.startswith("eyJ"):
                                    bearer_token = t
                                    print("  [OK] Token JWT extraído do Storage do navegador!")
                                    break
                except Exception as e:
                    print(f"  [Aviso] Falha ao inspecionar Web Storage: {e}")

            captured_cookies = context.cookies()
            print(f"Sessão validada. Total de cookies capturados: {len(captured_cookies)}")

        finally:
            page.wait_for_timeout(2_000)
            context.close()
            browser.close()

    if bearer_token:
        auth_headers["Authorization"] = f"Bearer {bearer_token}"
    
    session = requests.Session()
    session.headers.update(auth_headers)
    for c in captured_cookies:
        session.cookies.set(c["name"], c["value"], domain=c.get("domain", ""))

    return {
        "session": session,
        "token": bearer_token,
        "headers": auth_headers,
    }


# ─────────────────────────────────────────────
# 4. Requisições na API de Picking com Rate Limiter
# ─────────────────────────────────────────────

def consultar_api_picking(
    session_data: Dict[str, Any],
    nota: str,
    pdv: str,
    timeout_segundos: int = 30,
) -> Optional[List[Dict[str, Any]]]:
    """
    Consulta o endpoint de picking para uma nota fiscal e PDV específicos.
    Endpoint: /invoice-number/{nota}/store/{pdv}
    """
    url = f"{API_BASE_URL}/invoice-number/{nota}/store/{pdv}"
    session: requests.Session = session_data["session"]
    headers = session_data["headers"]

    max_tentativas = 3
    for tentativa in range(1, max_tentativas + 1):
        try:
            resp = session.get(url, headers=headers, timeout=timeout_segundos)
            
            if resp.status_code == 429:
                wait_time = 10 * tentativa
                print(f"  [Rate Limit 429] Aguardando {wait_time}s para tentar novamente...")
                time.sleep(wait_time)
                continue

            if resp.status_code == 200:
                try:
                    data = resp.json()
                    if isinstance(data, list):
                        return data
                    elif isinstance(data, dict):
                        if "items" in data and isinstance(data["items"], list):
                            return data["items"]
                        if "data" in data and isinstance(data["data"], list):
                            return data["data"]
                        if "volumes" in data and isinstance(data["volumes"], list):
                            itens_achatados = []
                            for vol in data["volumes"]:
                                vol_id = vol.get("volume") or vol.get("volumeIdentifier")
                                for it in vol.get("items", [vol]):
                                    it_copy = dict(it)
                                    if "volume" not in it_copy and vol_id:
                                        it_copy["volume"] = vol_id
                                    itens_achatados.append(it_copy)
                            return itens_achatados
                        return [data]
                except Exception as ex_json:
                    print(f"  [Erro JSON] Resposta inválida para Nota {nota} / PDV {pdv}: {ex_json}")
                    return None

            elif resp.status_code == 404:
                print(f"  [404 Não Encontrado] Nota {nota} / PDV {pdv} sem registros de picking.")
                return None
            else:
                print(f"  [Aviso] Nota {nota} / PDV {pdv} retornou HTTP {resp.status_code}: {resp.text[:120]}")

        except requests.RequestException as e:
            print(f"  [Erro de Rede] Tentativa {tentativa}/{max_tentativas} para Nota {nota} / PDV {pdv}: {e}")
            time.sleep(3)

    return None


# ─────────────────────────────────────────────
# 5. Geração de Planilha Individual (Padrão op_piking.xlsx)
# ─────────────────────────────────────────────

def formatar_item_para_colunas(item: Dict[str, Any]) -> Dict[str, Any]:
    """Mapeia os campos do JSON para as 9 colunas exatas do template op_piking.xlsx."""
    return {
        "volume": item.get("volume") or item.get("volumeIdentifier"),
        "itemNumber": str(item.get("itemNumber", "")).zfill(6) if item.get("itemNumber") is not None else "",
        "handlingUnitCode": str(item.get("handlingUnitCode", "")),
        "quantity": item.get("quantity") if item.get("quantity") is not None else 0,
        "batchCode": str(item.get("batchCode", "")),
        "volumeIdentifier": item.get("volumeIdentifier") or item.get("volume"),
        "materialCode": str(item.get("materialCode", "")).zfill(18) if item.get("materialCode") is not None else "",
        "materialDescription": str(item.get("materialDescription", "")),
        "expirationDate": str(item.get("expirationDate", "")),
    }


def gerar_planilha_op_picking_individual(
    itens: List[Dict[str, Any]],
    caminho_saida: Path | str,
) -> Path:
    """
    Gera a planilha XLSX individual de uma nota no formato rigoroso de op_piking.xlsx.
    Aba 'data', 9 colunas, larguras de coluna e tipos devidamente preservados.
    """
    caminho = Path(caminho_saida)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"

    # Escreve a linha de cabeçalho
    for col_idx, col_name in enumerate(COLUNAS_OP_PIKING, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(name="Calibri", size=11, bold=False)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 13.0

    # Escreve os registros de dados
    for row_idx, item in enumerate(itens, start=2):
        row_dict = formatar_item_para_colunas(item)
        for col_idx, col_name in enumerate(COLUNAS_OP_PIKING, start=1):
            val = row_dict.get(col_name)
            cell = ws.cell(row=row_idx, column=col_idx)
            
            if col_name == "quantity":
                try:
                    cell.value = float(val) if "." in str(val) else int(val)
                except Exception:
                    cell.value = val
            elif col_name in ("itemNumber", "handlingUnitCode", "materialCode", "expirationDate", "batchCode"):
                cell.number_format = "@"
                cell.value = str(val) if val is not None else ""
            else:
                cell.value = val

    wb.save(caminho)
    return caminho


def salvar_json_consolidado(
    dados_consolidados: Dict[str, Any],
    caminho_saida: Path | str,
) -> Path:
    """Função utilitária preservada para futura integração via JSON."""
    caminho = Path(caminho_saida)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(dados_consolidados, f, ensure_ascii=False, indent=2)

    return caminho


# ─────────────────────────────────────────────
# 6. Execução Principal (Orquestração Semanal)
# ─────────────────────────────────────────────

def parse_args():
    parser = argparse.ArgumentParser(description="Extração Semanal de Planilhas OP Picking por Nota (Branch T:\\13-Publica)")
    parser.add_argument("--relatorio", type=str, nargs="*", default=None, help="Caminho(s) do(s) relatório(s) do Matheus (.xls ou .xlsx)")
    parser.add_argument("--pasta-xmls", type=str, default=None, help="Pasta de XML BAIXADOS da semana (se não informado, detecta em T:\\13-Publica)")
    parser.add_argument("--nota", type=str, default=None, help="Número de nota fiscal única para teste")
    parser.add_argument("--pdv", type=str, default=None, help="Código do PDV / Store único para teste")
    parser.add_argument("--nosso-pedido", type=str, default="", help="Nosso Pedido (usado em modo unitário de teste)")
    parser.add_argument("--delay", type=float, default=5.0, help="Intervalo mínimo em segundos entre chamadas (padrão: 5.0s)")
    parser.add_argument("--gerar-json", action="store_true", help="Opcional: também salvar o grande JSON consolidado em paralelo")
    parser.add_argument("--data", type=str, default=None, help="Data de referência (DD/MM/YYYY) para compatibilidade com o Scheduler")
    parser.add_argument("--dias", type=int, default=0, help="Dias retroativos para compatibilidade com o Scheduler")
    parser.add_argument("--unidade", type=str, default=None, help="Unidade regional (BH ou UDI) para compatibilidade com o Scheduler")
    parser.add_argument("--headless", action="store_true", help="Executar em modo headless / não-interativo")
    return parser.parse_args()


@log_execution(modulo="XML Picking Semanal")
def main() -> None:
    args = parse_args()
    print("==================================================================")
    print("   SOFI — EXTRAÇÃO DE PLANILHAS OP PICKING POR NOTA FISCAL (T:13) ")
    print("==================================================================")

    notas_para_processar: List[Dict[str, Any]] = []
    pasta_semana_relatorios: Optional[Path] = None

    # 1. Modo direto (Teste unitário)
    if args.nota and args.pdv:
        notas_para_processar.append({
            "numero_nota": args.nota.strip(),
            "nosso_pedido": args.nosso_pedido.strip() if args.nosso_pedido else "",
            "pdv": args.pdv.strip(),
            "nome_loja": "",
            "data_emissao": "",
            "unidade": "",
        })
        print(f"Modo direto ativado: Nota={args.nota}, PDV={args.pdv}, Nosso Pedido={args.nosso_pedido}")

    # 2. Modo com relatórios informados via CLI
    elif args.relatorio:
        for r_path in args.relatorio:
            p_rel = Path(r_path)
            notas_para_processar.extend(extrair_notas_relatorio_matheus(p_rel))
            if not pasta_semana_relatorios:
                pasta_semana_relatorios = p_rel.parent

    # 3. Modo automático: busca os relatórios do Matheus da ÚLTIMA semana em T:\13-Publica
    else:
        relatorios_matheus, pasta_semana_relatorios = localizar_ultimos_relatorios_matheus()
        if relatorios_matheus:
            print(f"\n{len(relatorios_matheus)} relatório(s) do Matheus detectado(s) da última semana:")
            for r in relatorios_matheus:
                print(f" - {r.name}")
                notas_para_processar.extend(extrair_notas_relatorio_matheus(r))
        else:
            print("\nNenhum relatório do Matheus encontrado automaticamente.")
            if args.headless or not sys.stdin.isatty():
                print("[AVISO] Ambiente não interativo / scheduler: nenhum relatório encontrado. Encerrando.")
                sys.exit(0)
            rel_manual = input("Digite o caminho do relatório do Matheus: ").strip()
            if rel_manual:
                p_man = Path(rel_manual)
                notas_para_processar.extend(extrair_notas_relatorio_matheus(p_man))
                pasta_semana_relatorios = p_man.parent
            else:
                print("Operação cancelada.")
                sys.exit(0)

    # Remove duplicidades preservando os dados completos
    notas_unicas: List[Dict[str, Any]] = []
    vistos = set()
    for n in notas_para_processar:
        chave = (n["numero_nota"], n["pdv"])
        if chave not in vistos:
            vistos.add(chave)
            notas_unicas.append(n)

    if not notas_unicas:
        print("Nenhuma nota fiscal para consultar. Finalizando.")
        sys.exit(0)

    # 4. Localiza a pasta correspondente de XML BAIXADOS em T:\13-Publica (organizada por PDV)
    if args.pasta_xmls:
        pasta_xmls_semana = Path(args.pasta_xmls)
    else:
        pasta_xmls_semana = localizar_pasta_xmls_semana(pasta_semana_relatorios)

    print(f"\nPasta da semana em XML BAIXADOS: {pasta_xmls_semana}")
    print(f"Total de notas faturadas a consultar na semana: {len(notas_unicas)}")
    delay_base = max(5.0, args.delay)
    print(f"Controle de taxa ativo: delay mínimo de {delay_base:.1f}s (+ jitter) entre requisições.")

    # 5. Autenticação no Extranet e obtenção da sessão
    session_data = autenticar_extranet_e_obter_sessao()

    # 6. Loop de consultas na API de Picking e Geração Individual por Nota
    planilhas_geradas: List[Path] = []
    notas_resultado_json: List[Dict[str, Any]] = []
    itens_consolidados_json: List[Dict[str, Any]] = []
    total = len(notas_unicas)

    print("\n--- Iniciando consultas de Picking e gravação das planilhas por PDV ---")
    for idx, nota_info in enumerate(notas_unicas, start=1):
        n_nota = nota_info["numero_nota"]
        n_pdv = nota_info["pdv"]
        nosso_ped = nota_info.get("nosso_pedido", "")

        print(f"\n[{idx}/{total}] Nota: {n_nota} | Nosso Pedido: {nosso_ped} | PDV: {n_pdv}...")
        itens_picking = consultar_api_picking(session_data, n_nota, n_pdv) or []

        # Determina o caminho da planilha: <pasta_semana>/<PDV>/<PDV>_<Nº_DOC>.xlsx
        caminho_xlsx = obter_caminho_planilha_picking(pasta_xmls_semana, n_pdv, n_nota)

        # Gera a planilha individual da nota (aba 'data', 9 colunas no padrão op_piking.xlsx)
        gerar_planilha_op_picking_individual(itens_picking, caminho_xlsx)
        planilhas_geradas.append(caminho_xlsx)

        if itens_picking:
            print(f"  [OK] {len(itens_picking)} itens de picking salvos em: {caminho_xlsx.parent.name}\\{caminho_xlsx.name}")
        else:
            print(f"  [INFO] 0 itens de picking para Nota {n_nota} (planilha criada: {caminho_xlsx.parent.name}\\{caminho_xlsx.name})")

        # Estrutura auxiliar para o JSON (preservada para integração futura)
        if args.gerar_json:
            nota_info_copy = dict(nota_info)
            nota_info_copy["itens_picking"] = itens_picking
            nota_info_copy["total_itens_picking"] = len(itens_picking)
            nota_info_copy["arquivo_xlsx"] = str(caminho_xlsx)
            notas_resultado_json.append(nota_info_copy)
            for it in itens_picking:
                itens_consolidados_json.append({
                    "numero_nota": n_nota,
                    "nosso_pedido": nosso_ped,
                    "pdv": n_pdv,
                    "nome_loja": nota_info.get("nome_loja", ""),
                    "data_emissao": nota_info.get("data_emissao", ""),
                    "unidade": nota_info.get("unidade", ""),
                    "volume": it.get("volume") or it.get("volumeIdentifier"),
                    "itemNumber": str(it.get("itemNumber", "")).zfill(6) if it.get("itemNumber") is not None else "",
                    "handlingUnitCode": str(it.get("handlingUnitCode", "")),
                    "quantity": it.get("quantity") if it.get("quantity") is not None else 0,
                    "batchCode": str(it.get("batchCode", "")),
                    "volumeIdentifier": it.get("volumeIdentifier") or it.get("volume"),
                    "materialCode": str(it.get("materialCode", "")).zfill(18) if it.get("materialCode") is not None else "",
                    "materialDescription": str(it.get("materialDescription", "")),
                    "expirationDate": str(it.get("expirationDate", "")),
                })

        # Delay seguro entre requisições para proteger a API
        if idx < total:
            tempo_espera = delay_base + random.uniform(0.5, 1.5)
            print(f"  Aguardando {tempo_espera:.1f}s antes da próxima nota...")
            time.sleep(tempo_espera)

    # 7. Exportação opcional de JSON se solicitada via flag
    if args.gerar_json:
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        caminho_json = pasta_xmls_semana / f"picking_semanal_{timestamp_str}.json"
        salvar_json_consolidado({
            "metadata": {
                "gerado_em": datetime.now().isoformat(),
                "total_notas": len(notas_unicas),
                "total_planilhas_geradas": len(planilhas_geradas),
            },
            "notas": notas_resultado_json,
            "itens_consolidados": itens_consolidados_json,
        }, caminho_json)
        print(f"\n[OK] JSON auxiliar consolidado salvo em: {caminho_json}")

    print("\n==================================================================")
    print("Extração semanal finalizada com sucesso!")
    print(f"Total de planilhas de picking geradas: {len(planilhas_geradas)}")
    print(f"Estrutura de salvamento: <pasta_semana>\\<PDV>\\<PDV>_<Nº_DOC>.xlsx")
    print(f"Pasta da semana: {pasta_xmls_semana}")
    print("==================================================================")
    sys.exit(0)


if __name__ == "__main__":
    main()
