"""
gerador_relatorio_detalhado.py — Geração de relatórios de Nota Fiscal Detalhado por PDV (VD Captação).

Padrão de Nomenclatura Canônica:
    13-PUB.XML.VD.CAPTACAO.<PDV>.NOTA FISCAL DETALHADO.<SEMANA>.xlsx
    Exemplo: 13-PUB.XML.VD.CAPTACAO.20844.NOTA FISCAL DETALHADO.Semana 01-08 a 07-08.xlsx

Estrutura de Pastas:
    T:\\13-Publica\\XMLS CÁLAMO\\RELATÓRIOS DETALHADOS\\<ANO>\\<MÊS>\\<SEMANA>\\
    (com fallback para pasta local 'Relatórios Cálamo Detalhado')

Colunas da Planilha (Linha 1, Colunas A a E):
    A: Numero da Nota
    B: Codigo (PDV)
    C: Emissao (Data)
    D: SKU (Codigo SKU)
    E: Quantidade
"""

import os
import sys
import argparse
from pathlib import Path
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Localiza dinamicamente a raiz do projeto SOFI
_ROOT = Path(__file__).resolve().parent
while not (_ROOT / "core").exists() and _ROOT != _ROOT.parent:
    _ROOT = _ROOT.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

try:
    from core.sofi_env import load_sofi_env
    load_sofi_env(department="financeiro")
except Exception:
    pass

try:
    from core.database.logger import log_execution
except Exception:
    def log_execution(modulo="XML Detalhado"):
        def decorator(func):
            def wrapper(*args, **kwargs):
                return func(*args, **kwargs)
            return wrapper
        return decorator

from departments.financeiro.modules.xml.leitor_xml import extrair_itens_xml

# Configuração Padrão
PDVS_PADRAO_CAPTACAO = ['20844', '13793']
REDE_BASE_T = Path(r"T:\13-Publica\XMLS CÁLAMO\RELATÓRIOS DETALHADOS")
LOCAL_BASE_FALLBACK = Path(__file__).resolve().parent / "Relatórios Cálamo Detalhado"


def obter_diretorio_base_detalhado() -> Path:
    """Retorna o diretório base de saída na rede T: ou o fallback local se inacessível."""
    try:
        if REDE_BASE_T.parent.exists():
            return REDE_BASE_T
    except Exception:
        pass
    return LOCAL_BASE_FALLBACK


def salvar_planilha_pdv(pdv: str, itens: List[Dict[str, Any]], caminho_saida: str) -> bool:
    """Gera e formata a planilha Excel individual para o PDV com seus respectivos SKUs."""
    if not itens:
        return False

    os.makedirs(os.path.dirname(caminho_saida), exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nota fiscal Detalhado"

    # Cabeçalhos na Linha 1, Colunas A a E
    cabecalhos = ['Numero da Nota', 'Codigo', 'Emissao', 'SKU', 'Quantidade']
    ws.append(cabecalhos)

    # Estilos de cabeçalho
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, len(cabecalhos) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # Ordena os itens por Número da Nota e por SKU
    def chave_ordenacao(item):
        n_doc = item.get('n_doc', 0)
        try:
            n_doc_int = int(n_doc)
        except (ValueError, TypeError):
            n_doc_int = 0
        return (n_doc_int, str(item.get('sku', '')))

    itens_ordenados = sorted(itens, key=chave_ordenacao)

    # Escreve linhas de dados
    font_data = Font(name="Calibri", size=10)
    for row_idx, item in enumerate(itens_ordenados, start=2):
        ws.cell(row=row_idx, column=1, value=item.get('n_doc'))
        ws.cell(row=row_idx, column=2, value=str(item.get('pdv', pdv)))
        ws.cell(row=row_idx, column=3, value=item.get('dt_emissao'))
        ws.cell(row=row_idx, column=4, value=str(item.get('sku', '')))
        ws.cell(row=row_idx, column=5, value=item.get('quantidade'))

        # Aplica alinhamento e fontes
        ws.cell(row=row_idx, column=1).alignment = align_center
        ws.cell(row=row_idx, column=2).alignment = align_center
        ws.cell(row=row_idx, column=3).alignment = align_center
        ws.cell(row=row_idx, column=4).alignment = align_left
        ws.cell(row=row_idx, column=5).alignment = align_right

        for c in range(1, 6):
            cell = ws.cell(row=row_idx, column=c)
            cell.font = font_data
            cell.border = thin_border

    # Ajusta largura das colunas
    larguras = {1: 18, 2: 12, 3: 15, 4: 20, 5: 15}
    for col_idx, largura in larguras.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = largura

    # Congela painel na primeira linha
    ws.freeze_panes = "A2"

    wb.save(caminho_saida)
    wb.close()
    return True


def processar_relatorios_detalhados(
    lista_arquivos_xml: List[str],
    diretorio_base_saida: Optional[str] = None,
    ano_pasta: Optional[str] = None,
    mes_pasta: Optional[str] = None,
    semana_pasta: Optional[str] = None,
    pdvs_alvo: Optional[List[str]] = None,
) -> Dict[str, str]:
    """
    Processa os XMLs fornecidos, agrupa por PDV e gera as planilhas detalhadas para os PDVs alvo.
    
    Retorna:
        Dict mapeando pdv -> caminho_completo_do_arquivo_gerado
    """
    if pdvs_alvo is None:
        pdvs_alvo = PDVS_PADRAO_CAPTACAO

    pdvs_alvo_set = {str(p).strip() for p in pdvs_alvo}

    if not diretorio_base_saida:
        diretorio_base_saida = str(obter_diretorio_base_detalhado())

    # Agrupa itens por PDV
    itens_por_pdv: Dict[str, List[Dict[str, Any]]] = {}

    for caminho_xml in lista_arquivos_xml:
        if not os.path.exists(caminho_xml):
            continue
        itens = extrair_itens_xml(caminho_xml)
        for item in itens:
            pdv = str(item.get('pdv', '')).strip()
            if pdv in pdvs_alvo_set:
                if pdv not in itens_por_pdv:
                    itens_por_pdv[pdv] = []
                itens_por_pdv[pdv].append(item)

    if not itens_por_pdv:
        print(f"[Detalhado] Nenhum item encontrado para os PDVs alvo ({', '.join(pdvs_alvo)}).")
        return {}

    # Monta diretório de saída
    partes_caminho = [diretorio_base_saida]
    if ano_pasta:
        partes_caminho.append(ano_pasta)
    if mes_pasta:
        partes_caminho.append(mes_pasta)
    if semana_pasta:
        partes_caminho.append(semana_pasta)

    diretorio_destino = os.path.join(*partes_caminho)
    os.makedirs(diretorio_destino, exist_ok=True)

    sufixo_semana = semana_pasta if semana_pasta else "Semana_Geral"
    arquivos_gerados: Dict[str, str] = {}

    for pdv, lista_itens in itens_por_pdv.items():
        nome_arquivo = f"13-PUB.XML.VD.CAPTACAO.{pdv}.NOTA FISCAL DETALHADO.{sufixo_semana}.xlsx"
        caminho_arquivo = os.path.join(diretorio_destino, nome_arquivo)
        sucesso = salvar_planilha_pdv(pdv, lista_itens, caminho_arquivo)
        if sucesso:
            arquivos_gerados[pdv] = caminho_arquivo
            print(f" -> Relatório detalhado gerado para PDV {pdv} ({len(lista_itens)} itens): {caminho_arquivo}")

    return arquivos_gerados


@log_execution(modulo="XML Detalhado por PDV")
def main():
    parser = argparse.ArgumentParser(description="Gerador de Relatórios de Nota Fiscal Detalhado por PDV")
    parser.add_argument("--pasta", type=str, help="Pasta contendo os arquivos XML a processar")
    parser.add_argument("--ano", type=str, default="", help="Ano da pasta de saída")
    parser.add_argument("--mes", type=str, default="", help="Mês da pasta de saída (ex: 08 - Agosto)")
    parser.add_argument("--semana", type=str, default="", help="Semana da pasta de saída (ex: Semana 01-08 a 07-08)")
    parser.add_argument("--pdvs", type=str, default="20844,13793", help="Lista de PDVs separados por vírgula")
    parser.add_argument("--saida", type=str, default="", help="Diretório base de saída (opcional)")

    args = parser.parse_args()

    pdvs_lista = [p.strip() for p in args.pdvs.split(",") if p.strip()]

    arquivos_xml = []
    if args.pasta and os.path.exists(args.pasta):
        for root_dir, _, files in os.walk(args.pasta):
            for f in files:
                if f.lower().endswith('.xml'):
                    arquivos_xml.append(os.path.join(root_dir, f))
    else:
        print("Uso: informe a --pasta com os arquivos XML para processamento.")
        return

    print(f"Iniciando processamento detalhado de {len(arquivos_xml)} XMLs para os PDVs: {pdvs_lista}")
    gerados = processar_relatorios_detalhados(
        lista_arquivos_xml=arquivos_xml,
        diretorio_base_saida=args.saida if args.saida else None,
        ano_pasta=args.ano,
        mes_pasta=args.mes,
        semana_pasta=args.semana,
        pdvs_alvo=pdvs_lista
    )
    print(f"Finalizado. {len(gerados)} planilhas geradas com sucesso.")


if __name__ == '__main__':
    main()
