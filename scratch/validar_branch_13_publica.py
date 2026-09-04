import sys
from pathlib import Path
import openpyxl

# Adiciona raiz ao sys.path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from departments.financeiro.modules.xml.extrator_picking import (
    obter_caminho_planilha_picking,
    gerar_planilha_op_picking_individual,
    COLUNAS_OP_PIKING,
)

def run_tests():
    print("--- 1. Testando estrutura de caminhos por PDV (<PDV>_<Nº_DOC>.xlsx) ---")
    pasta_semana_teste = _ROOT / "scratch" / "XML_BAIXADOS_TESTE" / "2026" / "07 - Julho" / "Semana 06-07 a 12-07"
    
    pdv_teste = "20843"
    nota_teste = "576355"
    caminho_saida = obter_caminho_planilha_picking(pasta_semana_teste, pdv_teste, nota_teste)
    
    print(f"Caminho gerado: {caminho_saida}")
    assert caminho_saida.parent.name == pdv_teste, f"Subpasta deve ser o PDV {pdv_teste}"
    assert caminho_saida.name == f"{pdv_teste}_{nota_teste}.xlsx", f"Nome deve ser {pdv_teste}_{nota_teste}.xlsx"

    print("\n--- 2. Testando geração de planilha de picking no caminho gerado ---")
    mock_itens = [
        {
            "volume": "8/13",
            "itemNumber": "000010",
            "handlingUnitCode": "00003030000002283141",
            "quantity": 3,
            "batchCode": "Z207",
            "volumeIdentifier": "8/13",
            "materialCode": "000000000000059875",
            "materialDescription": "BOTIK AGUA MICLR FAC CIC/PANT/CER 200ml",
            "expirationDate": "20280731"
        }
    ]

    gerar_planilha_op_picking_individual(mock_itens, caminho_saida)
    assert caminho_saida.exists(), "A planilha deve ter sido criada no disco"

    print("\n--- 3. Validando estrutura de colunas com o template op_piking.xlsx ---")
    wb_gerado = openpyxl.load_workbook(caminho_saida)
    wb_template = openpyxl.load_workbook(_ROOT / "departments" / "financeiro" / "modules" / "xml" / "op_piking.xlsx")

    assert "data" in wb_gerado.sheetnames, "Aba deve ser 'data'"
    ws_g = wb_gerado["data"]
    ws_t = wb_template["data"]

    cols_gerado = [ws_g.cell(1, c).value for c in range(1, 10)]
    cols_template = [ws_t.cell(1, c).value for c in range(1, 10)]

    print("Colunas geradas :", cols_gerado)
    print("Colunas template:", cols_template)
    assert cols_gerado == cols_template, "Colunas devem ser 100% idênticas ao template"

    # Valida integridade das células
    assert ws_g.cell(2, 1).value == "8/13"
    assert ws_g.cell(2, 2).value == "000010"
    assert ws_g.cell(2, 4).value == 3
    assert ws_g.cell(2, 7).value == "000000000000059875"

    print("\n[OK] TESTES DA BRANCH T:\\13-Publica APROVADOS COM SUCESSO!")

if __name__ == "__main__":
    run_tests()
