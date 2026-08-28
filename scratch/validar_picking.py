import sys
from pathlib import Path
import openpyxl

# Adiciona raiz ao sys.path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from departments.financeiro.modules.xml.extrator_picking import (
    extrair_notas_de_planilha,
    gerar_planilha_op_picking,
    COLUNAS_OP_PIKING
)

def run_tests():
    # 1. Testa extração do relatório de exemplo existente
    path_rel = _ROOT / "departments" / "financeiro" / "modules" / "xml" / "Relatorio_Final_Calamo_22-05-2026.xlsx"
    notas_pdvs = extrair_notas_de_planilha(path_rel)
    print(f"Notas e PDVs extraídos do relatório ({len(notas_pdvs)} itens):", notas_pdvs[:5])

    # 2. Testa geração de XLSX com dados mockados
    mock_data = [
        {
            "volume": "8/13",
            "itemNumber": "000010",
            "handlingUnitCode": "00003030000002283141",
            "quantity": 3,
            "batchCode": "Z207",
            "volumeIdentifier": "8/13",
            "materialCode": "000000000000059875",
            "materialDescription": "BOTIK AGUA MICLR FAC CIC/PANT/CER 200ml",
            "expirationDate": "20280731"
        }
    ]
    caminho_teste = _ROOT / "scratch" / "teste_op_picking.xlsx"
    caminho_teste.parent.mkdir(parents=True, exist_ok=True)
    gerar_planilha_op_picking(mock_data, caminho_teste)

    # 3. Compara estrutura gerada com o template op_piking.xlsx
    wb_gerado = openpyxl.load_workbook(caminho_teste)
    wb_template = openpyxl.load_workbook(_ROOT / "departments" / "financeiro" / "modules" / "xml" / "op_piking.xlsx")

    cols_gerado = [wb_gerado["data"].cell(1, c).value for c in range(1, 10)]
    cols_template = [wb_template["data"].cell(1, c).value for c in range(1, 10)]

    print("Colunas geradas :", cols_gerado)
    print("Colunas template:", cols_template)
    assert cols_gerado == cols_template, "Colunas não coincidem!"
    print("\n[OK] TESTE DE ESTRUTURA E TEMPLATE APROVADO COM SUCESSO!")

if __name__ == "__main__":
    run_tests()
