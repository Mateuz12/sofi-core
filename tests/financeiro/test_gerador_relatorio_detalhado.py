import os
import sys
import unittest
from unittest.mock import patch
import tempfile
from pathlib import Path
import openpyxl

# Adiciona a raiz do projeto ao sys.path
_ROOT = Path(__file__).resolve().parent.parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from departments.financeiro.modules.xml.gerador_relatorio_detalhado import (
    salvar_planilha_pdv,
    processar_relatorios_detalhados,
    PDVS_PADRAO_CAPTACAO
)

class TestGeradorRelatorioDetalhado(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_salvar_planilha_pdv_layout_e_dados(self):
        caminho_arquivo = os.path.join(self.temp_dir.name, "teste_pdv.xlsx")
        itens = [
            {'n_doc': 1002, 'pdv': '20844', 'dt_emissao': '02/08/2025', 'sku': 'SKU-B', 'quantidade': 5},
            {'n_doc': 1001, 'pdv': '20844', 'dt_emissao': '01/08/2025', 'sku': 'SKU-A', 'quantidade': 10},
            {'n_doc': 1001, 'pdv': '20844', 'dt_emissao': '01/08/2025', 'sku': 'SKU-C', 'quantidade': 2.5},
        ]

        sucesso = salvar_planilha_pdv('20844', itens, caminho_arquivo)
        self.assertTrue(sucesso)
        self.assertTrue(os.path.exists(caminho_arquivo))

        # Valida conteúdo e estrutura da planilha
        wb = openpyxl.load_workbook(caminho_arquivo)
        sheet = wb.active

        # Cabeçalhos na Linha 1, Colunas A a E
        esperado_headers = ['Numero da Nota', 'Codigo', 'Emissao', 'SKU', 'Quantidade']
        headers_lidos = [sheet.cell(row=1, column=col).value for col in range(1, 6)]
        self.assertEqual(headers_lidos, esperado_headers)

        # Dados ordenados por Número da Nota e SKU
        # Linha 2 deve ser n_doc 1001, SKU-A
        self.assertEqual(sheet.cell(row=2, column=1).value, 1001)
        self.assertEqual(sheet.cell(row=2, column=2).value, '20844')
        self.assertEqual(sheet.cell(row=2, column=3).value, '01/08/2025')
        self.assertEqual(sheet.cell(row=2, column=4).value, 'SKU-A')
        self.assertEqual(sheet.cell(row=2, column=5).value, 10)

        # Linha 3 deve ser n_doc 1001, SKU-C
        self.assertEqual(sheet.cell(row=3, column=1).value, 1001)
        self.assertEqual(sheet.cell(row=3, column=4).value, 'SKU-C')
        self.assertEqual(sheet.cell(row=3, column=5).value, 2.5)

        # Linha 4 deve ser n_doc 1002, SKU-B
        self.assertEqual(sheet.cell(row=4, column=1).value, 1002)
        self.assertEqual(sheet.cell(row=4, column=4).value, 'SKU-B')
        self.assertEqual(sheet.cell(row=4, column=5).value, 5)

        wb.close()

    def test_processar_relatorios_detalhados_nomenclatura_e_pastas(self):
        # Cria XMLs de teste para 20844 e 13793
        xml_20844 = """<?xml version="1.0" encoding="utf-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe">
  <NFe>
    <infNFe>
      <ide><nNF>5001</nNF><dhEmi>2025-08-01T10:00:00-03:00</dhEmi></ide>
      <dest><CNPJ>20798658000431</CNPJ></dest>
      <det nItem="1"><prod><cProd>SKU-UDI-1</cProd><qCom>15</qCom></prod></det>
    </infNFe>
  </NFe>
</nfeProc>"""

        xml_13793 = """<?xml version="1.0" encoding="utf-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe">
  <NFe>
    <infNFe>
      <ide><nNF>6001</nNF><dhEmi>2025-08-02T11:00:00-03:00</dhEmi></ide>
      <dest><CNPJ>03982400000878</CNPJ></dest>
      <det nItem="1"><prod><cProd>SKU-BH-1</cProd><qCom>20</qCom></prod></det>
    </infNFe>
  </NFe>
</nfeProc>"""

        p1 = os.path.join(self.temp_dir.name, "nota_udi.xml")
        p2 = os.path.join(self.temp_dir.name, "nota_bh.xml")
        with open(p1, "w", encoding="utf-8") as f: f.write(xml_20844)
        with open(p2, "w", encoding="utf-8") as f: f.write(xml_13793)

        saida_base = os.path.join(self.temp_dir.name, "RELATORIOS_DETALHADOS")
        semana_str = "Semana 01-08 a 07-08"

        # Mock de mapeamento para garantir PDVs 20844 e 13793
        mapa_mock = {
            '20798658000431': {'PDV': '20844', 'Nomenclatura': 'UDI VD CAPTACAO', 'Canal': 'VD', 'Unidade': 'UDI'},
            '03982400000878': {'PDV': '13793', 'Nomenclatura': 'BH VD CAPTACAO', 'Canal': 'VD', 'Unidade': 'BH'}
        }

        with patch('departments.financeiro.modules.xml.leitor_xml.LOJAS_MAP_DINAMICO', mapa_mock):
            gerados = processar_relatorios_detalhados(
                lista_arquivos_xml=[p1, p2],
                diretorio_base_saida=saida_base,
                ano_pasta="2025",
                mes_pasta="08 - Agosto",
                semana_pasta=semana_str,
                pdvs_alvo=['20844', '13793']
            )

        self.assertIn('20844', gerados)
        self.assertIn('13793', gerados)

        # Checa nome exato dos arquivos
        nome_esperado_20844 = f"13-PUB.XML.VD.CAPTACAO.20844.NOTA FISCAL DETALHADO.{semana_str}.xlsx"
        nome_esperado_13793 = f"13-PUB.XML.VD.CAPTACAO.13793.NOTA FISCAL DETALHADO.{semana_str}.xlsx"

        self.assertEqual(os.path.basename(gerados['20844']), nome_esperado_20844)
        self.assertEqual(os.path.basename(gerados['13793']), nome_esperado_13793)

        # Checa se os arquivos foram gravados no disco
        self.assertTrue(os.path.exists(gerados['20844']))
        self.assertTrue(os.path.exists(gerados['13793']))

if __name__ == '__main__':
    unittest.main()
